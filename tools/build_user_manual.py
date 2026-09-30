from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\Proyectos\Lethe")
OUTPUT = ROOT / "entrega" / "Lethe-Final" / "Manual_de_usuario_Lethe.docx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17364D"
PALE = "EEF4F7"
LIGHT = "F7F9FA"
GRAY = RGBColor(75, 82, 87)


def set_font(run, name="Aptos", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)


def set_cell_border(cell, color="D9D9D9", size="8"):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = borders.find(qn("w:" + edge))
        if item is None:
            item = OxmlElement("w:" + edge)
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + tag))
        if node is None:
            node = OxmlElement("w:" + tag)
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9.5, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
            cells[idx].width = Inches(widths[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[idx])
            set_cell_margins(cells[idx])
            if row_index % 2:
                set_cell_shading(cells[idx], LIGHT)
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet 2" if level else "List Bullet")
    p.add_run(text)
    return p


step_counter = 0


def reset_steps():
    global step_counter
    step_counter = 0


def step(doc, text):
    global step_counter
    step_counter += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.33)
    p.paragraph_format.first_line_indent = Inches(-0.33)
    p.add_run(f"{step_counter}.\t{text}")
    return p


def page_heading(doc, text):
    p = doc.add_heading(text, level=1)
    p.paragraph_format.page_break_before = True
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Página ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def remove_paragraph_borders(paragraph):
    props = paragraph._p.get_or_add_pPr()
    borders = props.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "between"):
        item = OxmlElement("w:" + edge)
        item.set(qn("w:val"), "nil")
        borders.append(item)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.85)
section.footer_distance = Inches(0.3)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.1
for name, size in (("Title", 29), ("Heading 1", 19), ("Heading 2", 13)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(7)

footer = section.footer.paragraphs[0]
add_page_number(footer)
for run in footer.runs:
    set_font(run, size=8, color=GRAY)

doc.add_paragraph("PROYECTO LETHE", style="Subtitle")
title = doc.add_paragraph("Manual básico de usuario", style="Title")
remove_paragraph_borders(title)
subtitle = doc.add_paragraph()
subtitle.add_run("Cómo conectar, abrir, usar y cerrar la memoria cifrada").bold = True
subtitle.runs[0].font.size = Pt(15)
doc.add_paragraph("Windows  |  Lethe 0.1  |  VeraCrypt")
doc.add_paragraph("29 de septiembre de 2026")
doc.add_paragraph()
intro = doc.add_paragraph()
intro.add_run("Para qué sirve este manual. ").bold = True
intro.add_run(
    "Explica el uso de Lethe desde el momento en que conectas la memoria. No necesitas "
    "saber de criptografía ni usar comandos para el uso diario."
)
state = doc.add_paragraph()
state.add_run("Estado actual de JORGITO. ").bold = True
state.add_run(
    "La memoria está preparada en E: con el release final y un contenedor vault.hc de 6 GiB. "
    "El uso diario comienza abriendo el panel; no hace falta repetir la preparación."
)
warning = doc.add_paragraph()
warning.add_run("Regla más importante. ").bold = True
warning.add_run(
    "El botón Restaurar memoria física borra todo el contenido de la memoria. No se usa para "
    "abrir archivos ni para corregir una contraseña."
)
doc.add_heading("Qué vas a aprender", level=2)
bullet(doc, "Preparar JORGITO una sola vez.")
bullet(doc, "Abrir los archivos señuelo o los archivos privados.")
bullet(doc, "Usar el modo seguro y el modo con escritura.")
bullet(doc, "Cerrar correctamente y retirar la memoria.")
bullet(doc, "Reconocer los mensajes más comunes y saber qué no hacer.")

page_heading(doc, "Preparación inicial")
doc.add_paragraph(
    "Esta sección se realiza una sola vez. Si la memoria ya contiene la carpeta Lethe y el "
    "archivo vault.hc, pasa directamente a Conectar y abrir Lethe."
)
doc.add_heading("Copiar el programa", level=2)
reset_steps()
step(doc, "Conecta JORGITO y abre el Explorador de archivos.")
step(doc, "Busca la unidad con la etiqueta JORGITO. La letra puede ser E:, F: u otra.")
step(doc, "En este equipo, abre D:\\Proyectos\\Lethe\\staging\\Lethe-Release-0.1-Final.")
step(doc, "Copia la carpeta completa a la memoria y cámbiale el nombre a Lethe.")
step(doc, "Comprueba que dentro de Lethe aparecen Abrir Lethe.cmd, Lethe.exe, Lethe.ps1 y la carpeta tools.")
doc.add_heading("Crear el contenedor cifrado", level=2)
reset_steps()
step(doc, "Abre Lethe\\tools\\VeraCrypt y ejecuta VeraCrypt Format.exe.")
step(doc, "Elige Crear un contenedor de archivos cifrado.")
step(doc, "Elige Volumen VeraCrypt oculto y después Modo normal.")
step(doc, "Como ubicación selecciona Lethe\\vault.hc dentro de JORGITO.")
step(doc, "Para esta memoria de 8 GB, usa un contenedor de 6 GiB.")
step(doc, "Crea el volumen exterior con una contraseña señuelo distinta de la contraseña real.")
step(doc, "Copia al volumen exterior algunos archivos normales que puedas mostrar.")
step(doc, "Continúa el asistente y crea el volumen oculto con la contraseña real.")
step(doc, "Guarda los archivos realmente privados solo dentro del volumen oculto.")
doc.add_paragraph(
    "No guardes las contraseñas en la memoria. Si olvidas una contraseña, Lethe no puede recuperarla."
)

page_heading(doc, "Conectar y abrir Lethe")
reset_steps()
step(doc, "Conecta JORGITO al equipo y espera a que Windows la muestre.")
step(doc, "Abre el Explorador de archivos y entra en JORGITO.")
step(doc, "Abre la carpeta Lethe.")
step(doc, "Haz doble clic en Abrir Lethe.cmd.")
step(doc, "Espera a que el panel muestre Bloqueado. Listo para abrir. Si no se actualiza, pulsa Actualizar estado.")
doc.add_heading("Lo que verás en el panel", level=2)
add_table(
    doc,
    ["Elemento", "Para qué sirve"],
    [
        ["Abrir en modo seguro", "Abre el volumen sin permitir cambios. Es la opción recomendada."],
        ["Habilitar la opción de escritura", "Permite activar el botón de escritura cuando realmente necesitas modificar archivos."],
        ["Abrir con escritura", "Abre el volumen para crear, editar o borrar archivos."],
        ["Actualizar estado", "Consulta el montaje real y vuelve a detectar la memoria si Windows cambia su letra."],
        ["Abrir archivos cifrados", "Abre la unidad virtual cuando Lethe confirma que está montada."],
        ["Desmontar volumen", "Cierra la unidad virtual y limpia la caché de VeraCrypt."],
        ["Restaurar memoria física", "Borra todo y devuelve JORGITO a una memoria normal."],
        ["Diagnóstico", "Muestra el estado, la letra virtual, la versión y la ubicación del contenedor."],
    ],
    [2.15, 4.25],
)
doc.add_paragraph(
    "La unidad cifrada se abre normalmente como L:. JORGITO conserva su propia letra física, "
    "por ejemplo E:. Son dos unidades diferentes y es normal ver ambas."
)

page_heading(doc, "Elegir qué contenido abrir")
doc.add_paragraph(
    "Pulsa Abrir en modo seguro. VeraCrypt mostrará su propio cuadro para pedir la contraseña. "
    "Lethe nunca ve ni guarda lo que escribes en ese cuadro."
)
add_table(
    doc,
    ["Lo que escribes", "Qué ocurre", "Qué debes hacer"],
    [
        ["Contraseña incorrecta", "No aparece ninguna unidad.", "Cierra el aviso o vuelve a intentarlo. No se borra nada."],
        ["Contraseña señuelo", "Se abre el volumen exterior en L:.", "Usa o muestra solamente los archivos normales preparados."],
        ["Contraseña real", "Se abre el volumen oculto en L:.", "Accede a los archivos privados."],
    ],
    [1.55, 2.15, 2.7],
)
doc.add_heading("Cómo saber que se abrió", level=2)
bullet(doc, "El diagnóstico indica Estado real: MOUNTED y unidad cifrada L:.")
bullet(doc, "En el Explorador de archivos aparece una nueva unidad L:.")
bullet(doc, "El botón Abrir archivos cifrados queda habilitado.")
bullet(doc, "Lethe no indica si L: es el volumen exterior o el oculto. Esto es intencional.")
doc.add_heading("Si la contraseña falla", level=2)
doc.add_paragraph(
    "Una contraseña incorrecta no activa contadores, no borra el contenedor y no modifica la "
    "memoria. Después de unos segundos, el panel puede mostrar No se abrió ningún volumen."
)

page_heading(doc, "Usar tus archivos")
doc.add_heading("Modo seguro", level=2)
doc.add_paragraph(
    "Usa Abrir en modo seguro cuando solo quieras leer, copiar hacia otro lugar o comprobar "
    "archivos. En este modo no puedes guardar cambios dentro de L:."
)
reset_steps()
step(doc, "Abre L: desde el Explorador de archivos.")
step(doc, "Abre o copia los archivos que necesites.")
step(doc, "Cuando termines, cierra los archivos y las ventanas que estén usando L:.")
doc.add_heading("Modo con escritura", level=2)
doc.add_paragraph(
    "Utilízalo solo cuando necesites crear, editar o borrar archivos dentro del volumen privado."
)
reset_steps()
step(doc, "Si hay un volumen abierto, pulsa Desmontar volumen primero.")
step(doc, "Marca Habilitar la opción de escritura.")
step(doc, "Pulsa Abrir con escritura y acepta la advertencia.")
step(doc, "Introduce la contraseña real en VeraCrypt.")
step(doc, "Trabaja con los archivos de L: y guarda normalmente.")
danger = doc.add_paragraph()
danger.add_run("No abras el volumen señuelo con escritura. ").bold = True
danger.add_run(
    "Escribir en el volumen exterior sin protección especial puede dañar el volumen oculto. "
    "En Lethe, reserva el modo con escritura para la contraseña real."
)

page_heading(doc, "Cerrar y retirar la memoria")
doc.add_paragraph(
    "Cerrar correctamente evita archivos dañados y elimina de la caché la contraseña utilizada."
)
reset_steps()
step(doc, "Guarda los cambios en todos los programas.")
step(doc, "Cierra los documentos y todas las ventanas que estén abiertas en L:.")
step(doc, "Vuelve al panel de Lethe y pulsa Desmontar volumen.")
step(doc, "Espera el mensaje Volumen cerrado y caché limpiada.")
step(doc, "Comprueba que L: desapareció del Explorador de archivos.")
step(doc, "Cierra el panel de Lethe.")
step(doc, "En Windows, usa Quitar hardware de forma segura para expulsar JORGITO.")
step(doc, "Retira físicamente la memoria cuando Windows indique que es seguro.")
doc.add_heading("Si el cierre normal falla", level=2)
doc.add_paragraph(
    "Cierra cualquier archivo o programa que esté usando L: e inténtalo de nuevo. Usa el cierre "
    "forzado solamente después de cerrar todo, porque puede interrumpir una escritura pendiente."
)
doc.add_heading("Lo que nunca debes hacer", level=2)
bullet(doc, "No retires JORGITO mientras L: siga visible.")
bullet(doc, "No apagues el equipo mientras se están guardando archivos.")
bullet(doc, "No cambies, muevas ni borres vault.hc desde la parte normal de la memoria.")
bullet(doc, "No compartas las dos contraseñas con la misma persona si necesitas negación plausible.")

page_heading(doc, "Problemas comunes")
add_table(
    doc,
    ["Mensaje o problema", "Significado", "Solución sencilla"],
    [
        ["Lethe no está preparado", "Falta vault.hc o la configuración no corresponde.", "Comprueba que vault.hc esté dentro de la carpeta Lethe."],
        ["No se encontró el contenedor cifrado", "Lethe no encuentra vault.hc.", "No crees otro encima. Busca la copia correcta o restaura tu respaldo."],
        ["L: está ocupada", "Otro disco o programa ya usa esa letra.", "Desconecta la otra unidad o pide ayuda para cambiar mount_letter."],
        ["No se abrió ningún volumen", "La contraseña fue incorrecta o se canceló VeraCrypt.", "Vuelve a pulsar Abrir en modo seguro e inténtalo otra vez."],
        ["El volumen continúa abierto", "Algún programa sigue usando L:.", "Cierra archivos y ventanas; después pulsa Desmontar volumen."],
        ["La letra de JORGITO cambió", "Windows asignó otra letra al reconectar.", "Mantén una sola memoria Lethe conectada y pulsa Actualizar estado."],
        ["Error controlado del panel", "Lethe impidió que una excepción cerrara la interfaz.", "Pulsa Actualizar estado. Si continúa, cierra el panel y abre de nuevo Abrir Lethe.cmd."],
        ["Windows pide permiso", "Una operación administrativa necesita UAC.", "Acepta solo si tú iniciaste una restauración intencional."],
        ["Se desconectó por accidente", "El volumen pudo quedar sin desmontar.", "Reconecta JORGITO, abre Lethe y revisa el estado antes de volver a trabajar."],
    ],
    [1.7, 2.05, 2.7],
)
doc.add_heading("Contraseñas olvidadas", level=2)
doc.add_paragraph(
    "No existe una contraseña maestra ni un botón de recuperación. Sin la contraseña correcta, "
    "los datos cifrados no pueden abrirse. Conserva una copia de seguridad separada de los datos importantes."
)

page_heading(doc, "Volver JORGITO a una memoria normal")
erase = doc.add_paragraph()
erase.add_run("Esta operación borra todo. ").bold = True
erase.add_run(
    "Elimina el contenedor, el programa y cualquier otro archivo de JORGITO. Solo debes usarla "
    "cuando ya tengas copia de lo necesario y quieras dejar de utilizar el cifrado."
)
reset_steps()
step(doc, "Cierra todos los archivos y pulsa Desmontar volumen.")
step(doc, "Comprueba que L: ya no aparece.")
step(doc, "Haz una copia de cualquier dato que quieras conservar.")
step(doc, "Pulsa Restaurar memoria física.")
step(doc, "Escribe la letra actual de JORGITO sin dos puntos, por ejemplo E.")
step(doc, "Revisa la etiqueta JORGITO, la capacidad y el número de serie mostrados.")
step(doc, "Acepta la primera advertencia solamente si el objetivo es correcto.")
step(doc, "Escribe exactamente la frase de confirmación que muestra el panel.")
step(doc, "Acepta el permiso de administrador de Windows y espera.")
step(doc, "Cuando Lethe confirme el resultado, el panel se cerrará.")
step(doc, "Retira y vuelve a conectar JORGITO para que Windows le asigne una letra.")
step(doc, "Comprueba que aparece como JORGITO, exFAT y sin la carpeta Lethe.")
doc.add_paragraph(
    "La restauración vuelve utilizable la memoria, pero no es un borrado forense garantizado para memoria flash."
)

page_heading(doc, "Guía rápida")
add_table(
    doc,
    ["Quiero", "Acción"],
    [
        ["Ver archivos sin cambiarlos", "Abrir Lethe.cmd > Abrir en modo seguro > escribir la contraseña en VeraCrypt."],
        ["Ver el contenido señuelo", "Usar la contraseña señuelo. La unidad virtual aparece como L:."],
        ["Ver los archivos privados", "Usar la contraseña real. La unidad virtual aparece como L:."],
        ["Modificar archivos privados", "Desmontar el volumen > habilitar escritura > Abrir con escritura > usar la contraseña real."],
        ["Terminar", "Cerrar archivos > Desmontar volumen > esperar caché limpiada > expulsar JORGITO."],
        ["Probar otra contraseña", "Cerrar el diálogo de VeraCrypt y volver a abrir en modo seguro."],
        ["Dejar la memoria normal", "Hacer copia > Restaurar memoria física > confirmar > reconectar."],
    ],
    [2.05, 4.4],
)
doc.add_heading("Reglas para recordar", level=2)
bullet(doc, "Para el uso diario, empieza siempre con Abrir en modo seguro.")
bullet(doc, "Usa escritura únicamente con la contraseña real.")
bullet(doc, "Cierra el volumen antes de retirar la memoria.")
bullet(doc, "Restaurar memoria física significa borrar todo.")
bullet(doc, "Lethe no puede recuperar contraseñas olvidadas.")
doc.add_paragraph()
doc.add_paragraph(
    "Si tienes dudas y hay datos importantes, no pulses Restaurar memoria física. Cierra Lethe, "
    "deja la memoria conectada y pide ayuda antes de continuar."
)

doc.save(OUTPUT)
print(OUTPUT)
