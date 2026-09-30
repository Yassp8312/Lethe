from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"D:\Proyectos\Lethe")
OUTPUT = ROOT / "entrega" / "Lethe-Final" / "Informe_tecnico_Lethe_Final.docx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "15324A"
BLUE = "276A8B"
PALE = "E9F1F5"
LIGHT = "F5F7F8"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(70, 78, 84)


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)


def set_cell_border(cell, color="B5C3CB", size="8"):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        item = borders.find(qn(tag))
        if item is None:
            item = OxmlElement(tag)
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Página ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def remove_paragraph_borders(paragraph):
    props = paragraph._p.get_or_add_pPr()
    borders = props.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "between"):
        item = OxmlElement("w:" + edge)
        item.set(qn("w:val"), "nil")
        borders.append(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
            set_cell_border(cells[idx])
            if row_index % 2:
                set_cell_shading(cells[idx], LIGHT)
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
        if widths:
            for idx, width in enumerate(widths):
                cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def bullet(doc, text, level=0):
    style = "List Bullet 2" if level else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


NUMBER_COUNTER = 0


def reset_numbering():
    global NUMBER_COUNTER
    NUMBER_COUNTER = 0


def numbered(doc, text):
    global NUMBER_COUNTER
    NUMBER_COUNTER += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run(f"{NUMBER_COUNTER}.\t{text}")
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.85)
section.footer_distance = Inches(0.3)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for style_name, size in (("Title", 30), ("Heading 1", 18), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)

header = section.header.paragraphs[0]
header.text = "LETHE  |  INFORME TÉCNICO"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor(70, 95, 110)
add_page_number(section.footer.paragraphs[0])
for run in section.footer.paragraphs[0].runs:
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY

doc.add_paragraph("PROYECTO LETHE", style="Subtitle")
title = doc.add_paragraph(style="Title")
title.add_run("Sistema de almacenamiento cifrado con negación plausible")
remove_paragraph_borders(title)
subtitle = doc.add_paragraph()
subtitle.add_run("Informe técnico y justificación de diseño").bold = True
subtitle.runs[0].font.size = Pt(16)
subtitle.paragraph_format.space_before = Pt(12)
doc.add_paragraph("Versión 0.1  |  Cierre final del proyecto")
doc.add_paragraph("Windows  |  VeraCrypt 1.26.29  |  Rust  |  Componentes gratuitos")
doc.add_paragraph("29 de septiembre de 2026")
doc.add_paragraph()
status = doc.add_paragraph()
status.add_run("Estado final: ").bold = True
status.add_run(
    "el release final fue validado con la memoria JORGITO. Las pruebas cubrieron contraseña "
    "incorrecta, volumen exterior, volumen oculto, montaje, desmontaje, reconexión y restauración física."
)
doc.add_paragraph(
    "Documento de entrega para comprender qué se construyó, cómo funciona, qué decisiones "
    "protegen los datos y qué se comprobó durante la prueba física final."
)
heading = doc.add_heading("Resumen ejecutivo", level=1)
heading.paragraph_format.page_break_before = True
doc.add_paragraph(
    "Lethe es un lanzador para Windows que utiliza VeraCrypt como motor criptográfico. "
    "El diseño busca tres resultados ante una contraseña: ninguna unidad para una clave "
    "incorrecta, el volumen exterior señuelo para una contraseña y el volumen oculto real "
    "para otra. Lethe no recibe ni almacena contraseñas; el diálogo oficial de VeraCrypt "
    "las solicita y decide qué volumen puede abrirse."
)
doc.add_paragraph(
    "La versión construida ofrece apertura de solo lectura, apertura con escritura mediante "
    "confirmación, desmontaje normal o forzado, actualización del estado real, acceso directo "
    "a la unidad cifrada, limpieza de caché, diagnóstico, preparación "
    "transaccional de paquetes locales y restauración controlada de discos virtuales. El "
    "Sprint 7 completa la restauración física con identidad exacta, doble confirmación y "
    "revalidación dentro del ayudante elevado."
)
doc.add_heading("Resultado de seguridad", level=2)
bullet(doc, "No se diseñó criptografía propia; se reutiliza el formato y la implementación de VeraCrypt.")
bullet(doc, "La contraseña no pasa por argumentos, archivos, variables de entorno ni registros de Lethe.")
bullet(doc, "La escritura está deshabilitada por defecto para reducir el riesgo de dañar el volumen oculto.")
bullet(doc, "Una contraseña incorrecta no cambia el contenedor y no activa borrado ni contadores.")
bullet(doc, "La restauración física exige coincidencia exacta y una frase específica; nunca elige otro disco.")

doc.add_heading("Alcance de esta entrega", level=2)
add_table(
    doc,
    ["Componente", "Estado", "Evidencia"],
    [
        ["Núcleo Rust", "Completo", "45 pruebas y clippy sin advertencias"],
        ["Integración VeraCrypt", "Completa", "Versión y firma verificadas localmente"],
        ["Panel Windows", "Completo", "Montaje, estado real, acceso, desmontaje y restauración"],
        ["Restauración virtual", "Completa", "VHDX restaurado a GPT/exFAT/JORGITO"],
        ["Ciclo cifrado", "Completo", "Volúmenes exterior y oculto montados manualmente"],
        ["Restauración USB física", "Completa", "JORGITO quedó saludable en E:, GPT/exFAT"],
    ],
    [1.6, 1.2, 3.8],
)
heading = doc.add_heading("Requisitos y límites", level=1)
heading.paragraph_format.page_break_before = True
doc.add_heading("Requisitos funcionales", level=2)
reset_numbering()
numbered(doc, "Abrir el mismo contenedor con el diálogo seguro de VeraCrypt.")
numbered(doc, "No montar ninguna unidad cuando la contraseña no sea válida.")
numbered(doc, "Abrir el volumen exterior con la contraseña señuelo.")
numbered(doc, "Abrir el volumen oculto con la contraseña real.")
numbered(doc, "Cerrar el volumen y limpiar la caché de contraseñas.")
numbered(doc, "Restaurar el medio a un formato convencional desde el programa.")

doc.add_heading("Supuestos", level=2)
bullet(doc, "Windows 10 1809 o posterior.")
bullet(doc, "VeraCrypt 1.26.29 o una versión compatible y firmada.")
bullet(doc, "El usuario conserva contraseñas fuertes y copias de seguridad separadas.")
bullet(doc, "El adversario puede observar o destruir el dispositivo, pero no controla el sistema mientras se abre el volumen.")

doc.add_heading("Límites explícitos", level=2)
doc.add_paragraph(
    "La negación plausible no demuestra que un volumen oculto exista ni garantiza que un "
    "adversario crea una explicación. La existencia del contenedor VeraCrypt es visible. "
    "El cifrado protege confidencialidad; no impide que alguien borre o corrompa el archivo."
)
doc.add_paragraph(
    "El formateo convencional de memoria flash no equivale a borrado forense. La nivelación "
    "de desgaste puede dejar bloques que Windows no puede sobrescribir directamente."
)

doc.add_heading("Arquitectura", level=1)
add_table(
    doc,
    ["Capa", "Responsabilidad", "Dato sensible"],
    [
        ["Panel PowerShell", "Acciones de usuario y mensajes públicos", "Ninguno"],
        ["Núcleo Rust", "Reglas, estados, rutas y validaciones", "Ninguno"],
        ["Adaptador VeraCrypt", "Localizar, validar e invocar el motor", "No recibe contraseña"],
        ["VeraCrypt", "Solicitar contraseña, derivar claves y montar", "Contraseña en su diálogo"],
        ["Ayudante de restauración", "Revalidar y restaurar VHD/VHDX o USB autorizada", "Identidad técnica"],
    ],
    [1.5, 3.4, 1.7],
)
doc.add_paragraph(
    "El flujo de montaje es panel → núcleo → adaptador → VeraCrypt. El cuadro de contraseña "
    "pertenece a VeraCrypt. El flujo de restauración usa un ayudante temporal elevado y "
    "solo acepta el objetivo previamente planificado y vuelve a comprobar su identidad."
)
heading = doc.add_heading("Tres resultados de acceso", level=1)
heading.paragraph_format.page_break_before = True
add_table(
    doc,
    ["Entrada", "Resultado observable", "Acción de Lethe"],
    [
        ["Contraseña incorrecta", "No aparece unidad", "Espera; no modifica datos"],
        ["Contraseña señuelo", "Se monta el volumen exterior", "No distingue el tipo de volumen"],
        ["Contraseña real", "Se monta el volumen oculto", "No distingue el tipo de volumen"],
    ],
    [1.7, 2.4, 2.5],
)
doc.add_paragraph(
    "La indistinguibilidad dentro de Lethe es deliberada. Registrar qué contraseña se usó o "
    "intentar identificar el volumen abierto crearía información adicional que debilitaría "
    "el objetivo de negación plausible."
)

doc.add_heading("Política de escritura", level=2)
doc.add_paragraph(
    "El modo seguro abre en solo lectura. La escritura requiere marcar una opción y aceptar "
    "una advertencia. Escribir en el volumen exterior sin la protección adecuada de VeraCrypt "
    "puede sobrescribir sectores que pertenecen al volumen oculto. Para tareas normales se "
    "recomienda consultar el señuelo en solo lectura y reservar escritura para el volumen real."
)

doc.add_heading("Decisiones de diseño principales", level=1)
add_table(
    doc,
    ["Decisión", "Justificación", "Consecuencia"],
    [
        ["VeraCrypt como motor", "Evita inventar criptografía", "Compatibilidad con volúmenes ocultos"],
        ["Contraseña fuera de Lethe", "Evita exposición en procesos y registros", "El diálogo pertenece a VeraCrypt"],
        ["Contenedor fijo", "Simplifica respaldo y reversión", "El cifrado es visible"],
        ["Solo lectura por defecto", "Protege el volumen oculto", "La escritura exige confirmación"],
        ["Preparación transaccional", "No mezcla entregas parciales", "Destinos existentes se rechazan"],
        ["VHDX antes de USB", "Aísla primero las pruebas destructivas", "La USB se prueba tras validar el flujo"],
        ["Identidad ligada al plan", "Detecta cambios entre pasos", "Cualquier diferencia cancela"],
        ["Registro mínimo local", "Ayuda a diagnosticar sin crear secretos", "No bloquea operaciones"],
    ],
    [1.6, 2.9, 2.1],
)
heading = doc.add_heading("Implementación por sprints", level=1)
heading.paragraph_format.page_break_before = True
for heading, body in [
    ("Sprint 0", "Se definieron alcance, amenazas, estados, arquitectura, restricciones y decisiones iniciales."),
    ("Sprint 1", "Se creó el núcleo Rust con tipos seguros, configuración validada, estados y errores públicos."),
    ("Sprint 2", "Se integró VeraCrypt, incluida la validación de versión y firma y comandos sin contraseña."),
    ("Sprint 3", "Se construyó el panel WinForms para abrir, cerrar, diagnosticar y controlar escritura."),
    ("Sprint 4", "Se añadió preparación transaccional local y el asistente oficial para crear vault.hc."),
    ("Sprint 5", "Se implementó la restauración VHD/VHDX con identidad, UAC, GPT, exFAT y verificación."),
    ("Sprint 6", "Se añadió auditoría sin secretos, inspección USB dirigida, pruebas de fallo y entrega documental."),
    ("Sprint 7", "Se implementó y validó el ciclo físico completo, incluida la restauración segura de JORGITO."),
    ("Cierre final", "Se corrigió la ruta del lanzador, se desacopló PowerShell de la USB y se sincronizó el release final."),
]:
    doc.add_heading(heading, level=2)
    doc.add_paragraph(body)

doc.add_heading("Restauración controlada", level=1)
doc.add_heading("Secuencia virtual implementada", level=2)
reset_numbering()
numbered(doc, "Crear un plan sin montar ni escribir la imagen.")
numbered(doc, "Capturar ruta canónica, tamaño y marcas de tiempo.")
numbered(doc, "Mostrar el objetivo y exigir confirmación.")
numbered(doc, "Copiar un ayudante temporal y solicitar elevación de Windows.")
numbered(doc, "Revalidar la imagen y obtener el disco desde Mount-DiskImage.")
numbered(doc, "Rechazar cualquier disco que no sea virtual, o sea de arranque o sistema.")
numbered(doc, "Crear GPT, una partición exFAT y la etiqueta JORGITO.")
numbered(doc, "Verificar el resultado y desmontar en un bloque de limpieza.")

doc.add_heading("Fallo seguro", level=2)
doc.add_paragraph(
    "Una confirmación incorrecta, una identidad modificada, un objetivo externo o un bus "
    "inesperado detienen la operación. El proceso no intenta adivinar otro objetivo ni "
    "corregir automáticamente una discrepancia."
)

doc.add_heading("Prueba de la memoria física", level=1)
doc.add_paragraph(
    "El Sprint 7 reutilizó la inspección de solo lectura y añadió la ejecución física. El "
    "comando recibe una letra y resuelve desde ella el volumen, la partición única y el "
    "disco; el ayudante elevado revalida capacidad, serial, bus y banderas antes de escribir."
)
add_table(
    doc,
    ["Campo", "Valor observado"],
    [
        ["Letra final", "E:"],
        ["Etiqueta", "JORGITO"],
        ["Sistema de archivos", "exFAT"],
        ["Disco y partición", "Disco 1, partición 1"],
        ["Capacidad", "8 011 120 640 bytes"],
        ["Bus", "USB"],
        ["Serial", "057907B76060"],
        ["Arranque y sistema", "No / No"],
        ["Resultado", "Restauración completada; volumen saludable"],
    ],
    [2.2, 4.4],
)
doc.add_paragraph(
    "La letra y el número de disco pueden cambiar al reconectar; por eso no forman la "
    "identidad persistente. El serial 057907B76060 permaneció estable durante todo el ciclo."
)

doc.add_heading("Incidentes corregidos", level=2)
bullet(doc, "Clear-Disk dejó el dispositivo como MBR vacío; el ayudante ahora convierte MBR a GPT y solo inicializa si está RAW.")
bullet(doc, "Reasignar la letra mientras el panel se ejecutaba desde la USB produjo un falso fallo; la restauración termina sin letra y solicita reconectar.")
bullet(doc, "La contraseña exterior y la oculta fueron introducidas directamente por el usuario en VeraCrypt; Lethe nunca las conoció.")
bullet(doc, "El lanzador entregaba una comilla residual cuando HomePath terminaba en barra; ahora usa una ruta normalizada y ejecuta PowerShell desde la carpeta temporal local.")
bullet(doc, "El panel usa el estado del controlador como fuente de verdad, permite actualizar manualmente y puede redetectar una única carpeta Lethe válida.")
bullet(doc, "El release final y la copia instalada en JORGITO quedaron sincronizados con el mismo panel y ejecutable.")

doc.add_heading("Registro técnico", level=1)
doc.add_paragraph(
    "El registro se guarda en %LOCALAPPDATA%\\Lethe\\audit.log, fuera de la memoria. Cada "
    "línea contiene una marca de tiempo y un evento estructurado. No se escriben contraseñas, "
    "argumentos, rutas, letras, etiquetas, números de serie ni contenido del usuario."
)
add_table(
    doc,
    ["Evento", "Uso"],
    [
        ["app_started", "Inicio del proceso"],
        ["app_stopped", "Cierre normal"],
        ["operation_rejected code=…", "Código estable para diagnóstico"],
    ],
    [2.7, 3.9],
)
doc.add_heading("Pruebas y evidencia", level=1)
add_table(
    doc,
    ["Área", "Prueba", "Resultado"],
    [
        ["Núcleo", "45 pruebas unitarias", "Aprobado"],
        ["Calidad", "clippy con advertencias como errores", "Aprobado"],
        ["Compilación", "Perfil Release", "Aprobado"],
        ["Restauración", "VHDX de laboratorio a GPT/exFAT", "Aprobado"],
        ["Identidad", "Cambio después del plan", "Cancelado"],
        ["Confirmación", "Frase incorrecta", "Cancelado"],
        ["Destino", "Imagen fuera del sandbox", "Rechazado"],
        ["Cifrado", "Exterior y oculto montados manualmente", "Aprobado"],
        ["USB", "Restauración física y reconexión", "Aprobado"],
        ["USB", "Etiqueta, serial o banderas inseguras", "Rechazado"],
        ["Aceptación", "Contraseña incorrecta, exterior y oculta", "Aprobado"],
        ["Aceptación", "Desmontaje, reconexión y release copiado", "Aprobado"],
    ],
    [1.3, 3.5, 1.8],
)

doc.add_heading("Interpretación", level=2)
doc.add_paragraph(
    "Las pruebas muestran que las barreras previstas están presentes y que la restauración "
    "virtual y física funciona en el entorno utilizado. No sustituyen una auditoría criptográfica de "
    "VeraCrypt. El binario de Lethe tampoco posee todavía una firma "
    "de código propia."
)

doc.add_heading("Manual operativo", level=1)
doc.add_heading("Abrir y cerrar", level=2)
reset_numbering()
numbered(doc, "Inicia el panel desde panel\\Lethe.ps1.")
numbered(doc, "Usa el modo seguro de solo lectura como opción habitual.")
numbered(doc, "Introduce la contraseña únicamente en el diálogo de VeraCrypt.")
numbered(doc, "Para modificar el volumen real, habilita escritura y confirma la advertencia.")
numbered(doc, "Cierra los archivos y pulsa Desmontar volumen antes de retirar la memoria.")

doc.add_heading("Contraseña incorrecta", level=2)
doc.add_paragraph(
    "No aparece una unidad y no se altera el contenedor. Puede cerrarse el diálogo o volver a "
    "intentarse. El proyecto no incluye recuperación de contraseñas ni borrado por intentos."
)
heading = doc.add_heading("Recuperación y resultado final", level=1)
heading.paragraph_format.page_break_before = True
doc.add_heading("Qué hacer ante un fallo", level=2)
bullet(doc, "Guarda y cierra archivos antes de intentar un desmontaje forzado.")
bullet(doc, "Si cambia la identidad, cancela y crea un plan nuevo.")
bullet(doc, "Si Windows rechaza la elevación, la restauración no se ejecuta.")
bullet(doc, "Consulta el código del registro local sin buscar contraseñas en él.")
bullet(doc, "No desconectes un objetivo mientras una restauración esté en curso.")

doc.add_heading("Resultado final de la USB física", level=2)
bullet(doc, "El contenedor de 6 GiB admitió un volumen exterior señuelo y un volumen oculto real.")
bullet(doc, "El panel montó y desmontó el volumen, y limpió la caché de VeraCrypt.")
bullet(doc, "La restauración exigió plan visible, frase específica y revalidación elevada.")
bullet(doc, "Después de la prueba de restauración, JORGITO se preparó otra vez con el release final y vault.hc.")
bullet(doc, "Windows presenta la memoria física como E: y el volumen seleccionado por la contraseña como L:.")
bullet(doc, "La restauración convencional no garantiza borrado forense de una memoria flash.")

doc.add_heading("Trabajo pendiente", level=1)
bullet(doc, "Firmar el ejecutable de distribución si se publica fuera del entorno de desarrollo.")
bullet(doc, "Solicitar una auditoría externa si el proyecto se usará frente a adversarios de alto nivel.")

doc.add_heading("Conclusión", level=1)
doc.add_paragraph(
    "Lethe cierra el proyecto con un flujo funcional y conservador: delega la "
    "criptografía a VeraCrypt, mantiene las contraseñas fuera del programa, controla la "
    "escritura y restaura discos virtuales o la USB física exacta. La entrega final quedó "
    "sincronizada en staging y JORGITO terminó preparada para el uso cifrado diario."
)

doc.add_heading("Archivos de referencia", level=1)
for item in [
    "README.md",
    "docs/ARQUITECTURA.md",
    "docs/DECISIONES.md",
    "docs/MODELO_DE_AMENAZAS.md",
    "docs/ESTADOS.md",
    "docs/RECUPERACION.md",
    "docs/PRUEBA_FINAL_USB.md",
    "docs/SPRINT_0.md a docs/SPRINT_7.md",
]:
    bullet(doc, item)

doc.save(OUTPUT)
print(OUTPUT)
